"""Word document generation for extraction results."""

import logging
import tempfile
from datetime import datetime
from pathlib import Path
from typing import List, Optional

logger = logging.getLogger("api")


async def generate_word_document(extractor, usage_report, output_dir: Path) -> Optional[Path]:
    """Generate Word document with requirements registry and usage report.

    Args:
        extractor: RequirementsExtractor instance with populated registry.
        usage_report: UsageReport instance.
        output_dir: Directory where the .docx file will be saved.

    Returns:
        Path to the generated file, or None if python-docx is unavailable.
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.warning("python-docx not installed, skipping Word generation")
        return None

    doc = Document()

    title = doc.add_heading("Реестр требований", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Дата создания: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph(f"Всего требований: {extractor.registry.total_requirements}")
    doc.add_paragraph(f"Всего разделов: {len(extractor.registry.sections)}")
    doc.add_paragraph("")

    doc.add_heading("Требования по разделам", 1)
    for section in extractor.registry.sections:
        doc.add_heading(section.full_title, 2)
        doc.add_paragraph(f"Страницы: {section.page_range}")
        doc.add_paragraph(f"Требований: {len(section.requirements)}")
        doc.add_paragraph("")

        if section.requirements:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "Требование", "Тип", "Приоритет"
            for req in section.requirements:
                row = table.add_row().cells
                row[0].text = req.id
                row[1].text = req.text
                row[2].text = req.type.value if hasattr(req.type, "value") else str(req.type or "")
                row[3].text = req.priority.value if hasattr(req.priority, "value") else str(req.priority or "")
            doc.add_paragraph("")

    doc.add_page_break()
    doc.add_heading("Отчет о затратах", 1)
    doc.add_paragraph(f"Всего вызовов API: {usage_report.calls_count}")
    doc.add_paragraph(f"Входных токенов: {usage_report.total_input_tokens:,}")
    doc.add_paragraph(f"Выходных токенов: {usage_report.total_output_tokens:,}")
    doc.add_paragraph(f"Всего токенов: {usage_report.total_tokens:,}")
    doc.add_paragraph("")

    cost_run = doc.add_paragraph().add_run(f"Общая стоимость: ${usage_report.total_cost:.4f} USD")
    cost_run.bold = True
    cost_run.font.size = Pt(14)
    cost_run.font.color.rgb = RGBColor(0, 128, 0)
    doc.add_paragraph("")

    doc.add_heading("Детализация по разделам", 2)
    stats = usage_report.get_stats_by_section()
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Раздел", "Вызовов", "Токенов", "Стоимость"
    for sec_name, sec_stats in sorted(stats.items()):
        row = tbl.add_row().cells
        row[0].text = sec_name
        row[1].text = str(sec_stats["calls"])
        row[2].text = f"{sec_stats['input_tokens'] + sec_stats['output_tokens']:,}"
        row[3].text = f"${sec_stats['cost']:.4f}"

    word_path = output_dir / "requirements_report.docx"
    doc.save(word_path)
    logger.info(f"Word document generated: {word_path}")
    return word_path


def generate_word_from_db(document, sections: List, all_requirements: List, metrics) -> Optional[str]:
    """Generate Word document from database objects and return path to temp file.

    Args:
        document: DB Document ORM object.
        sections: List of DB Section ORM objects.
        all_requirements: List of DB Requirement ORM objects.
        metrics: DB CoverageMetrics ORM object or None.

    Returns:
        Path to the generated .docx temp file, or None if python-docx is unavailable.
    """
    try:
        from docx import Document as DocxDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.warning("python-docx not installed, skipping Word generation")
        return None

    doc = DocxDocument()

    title = doc.add_heading("Реестр требований", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Документ: {document.filename}")
    doc.add_paragraph(f"Дата обработки: {document.uploaded_at.strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph(f"Всего требований: {len(all_requirements)}")
    doc.add_paragraph(f"Всего разделов: {len(sections)}")
    doc.add_paragraph("")

    def _add_requirements_table(doc, reqs):
        table = doc.add_table(rows=1, cols=6)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "ID"
        hdr[1].text = "Требование"
        hdr[2].text = "Тип"
        hdr[3].text = "Приоритет"
        hdr[4].text = "Страница"
        hdr[5].text = "Статус"
        for req in reqs:
            row = table.add_row().cells
            row[0].text = req.requirement_id or ""
            # Show human-edited text if available, otherwise original AI text
            row[1].text = req.human_edited or req.text or ""
            row[2].text = req.type or ""
            row[3].text = req.priority or ""
            row[4].text = str(req.page_number or "")
            row[5].text = req.status or "pending"
        doc.add_paragraph("")

    doc.add_heading("Требования по разделам", 1)
    assigned_ids = set()
    for section in sections:
        section_reqs = [r for r in all_requirements if r.section_id == section.id]
        assigned_ids.update(r.id for r in section_reqs)

        doc.add_heading(section.title, 2)
        doc.add_paragraph(f"Страницы: {section.page_start} - {section.page_end}")
        doc.add_paragraph(f"Требований: {len(section_reqs)}")
        doc.add_paragraph("")

        if section_reqs:
            _add_requirements_table(doc, section_reqs)

    # Requirements not linked to any section
    orphans = [r for r in all_requirements if r.id not in assigned_ids]
    if orphans:
        doc.add_heading("Требования без раздела", 2)
        doc.add_paragraph(f"Требований: {len(orphans)}")
        doc.add_paragraph("")
        _add_requirements_table(doc, orphans)

    if metrics:
        doc.add_page_break()
        doc.add_heading("Метрики покрытия", 1)
        doc.add_paragraph(f"Всего страниц: {metrics.total_pages}")
        doc.add_paragraph(f"Обработано страниц: {metrics.processed_pages}")
        doc.add_paragraph(f"Пропущено страниц: {len(metrics.skipped_pages) if metrics.skipped_pages else 0}")
        doc.add_paragraph(f"Покрытие: {metrics.coverage_percent:.1f}%")
        doc.add_paragraph("")

        if metrics.requirements_by_type:
            doc.add_heading("Требования по типам", 2)
            table = doc.add_table(rows=1, cols=2)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text = "Тип", "Количество"
            for req_type, count in metrics.requirements_by_type.items():
                row = table.add_row().cells
                row[0].text = req_type
                row[1].text = str(count)

    tmp = tempfile.NamedTemporaryFile(mode="wb", suffix=".docx", delete=False)
    doc.save(tmp.name)
    tmp.close()
    logger.info(f"[EXPORT] Word document generated: {tmp.name}")
    return tmp.name
